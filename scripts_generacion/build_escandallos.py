#!/usr/bin/env python3
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

BLUE = Font(color="0000FF")
BLACK = Font(color="000000")
BOLD = Font(bold=True)
TITLE = Font(bold=True, size=14, color="1E3A5F")
HEAD = Font(bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor="1E3A5F")
YELLOW = PatternFill("solid", fgColor="FFFF00")
GREY = PatternFill("solid", fgColor="F1EDE0")
WRAP = Alignment(wrap_text=True, vertical="top")
THIN = Side(style="thin", color="D8D0BC")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

INGREDIENTES = {
    'papa':       dict(nombre='Papa negra canaria',       unidad='kg', precio=1.85, merma=8),
    'pim_rojo':   dict(nombre='Pimiento rojo',             unidad='kg', precio=3.05, merma=15),
    'pim_verde':  dict(nombre='Pimiento verde',            unidad='kg', precio=2.40, merma=15),
    'aceite':     dict(nombre='Aceite de oliva',           unidad='l',  precio=6.50, merma=0),
    'ajo':        dict(nombre='Ajo',                       unidad='kg', precio=8.00, merma=12),
    'comino':     dict(nombre='Comino molido',             unidad='kg', precio=14.00, merma=0),
    'sal':        dict(nombre='Sal gruesa',                unidad='kg', precio=0.80, merma=0),
    'vinagre':    dict(nombre='Vinagre de vino',           unidad='l',  precio=1.50, merma=0),
    'pimenton':   dict(nombre='Pimentón',                  unidad='kg', precio=12.00, merma=0),
    'ternera':    dict(nombre='Falda de ternera',          unidad='kg', precio=9.80, merma=20),
    'garbanzo':   dict(nombre='Garbanzo cocido',           unidad='kg', precio=2.20, merma=0),
    'cebolla':    dict(nombre='Cebolla',                   unidad='kg', precio=1.10, merma=10),
    'tomate':     dict(nombre='Tomate',                    unidad='kg', precio=1.95, merma=10),
    'vieja':      dict(nombre='Vieja (pescado local)',     unidad='kg', precio=12.50, merma=35),
    'gofio':      dict(nombre='Gofio de millo',            unidad='kg', precio=2.80, merma=0),
    'queso':      dict(nombre='Queso canario semicurado',  unidad='kg', precio=11.00, merma=5),
    'platano':    dict(nombre='Plátano de Canarias',       unidad='kg', precio=2.30, merma=30),
    'miel_palma': dict(nombre='Miel de palma',             unidad='l',  precio=15.00, merma=0),
    'vino':       dict(nombre='Vino blanco de cocina',     unidad='l',  precio=2.80, merma=0),
}

ALERGENOS_POR_INGREDIENTE = {
    'vinagre': {'Sulfitos': 'Contiene'},
    'vieja': {'Pescado': 'Contiene'},
    'gofio': {'Cereales con gluten': 'Pendiente de validar'},
    'queso': {'Leche y lactosa': 'Contiene'},
    'vino': {'Sulfitos': 'Contiene'},
}

LISTA_14_ALERGENOS = ['Cereales con gluten','Crustáceos','Huevos','Pescado','Cacahuetes','Soja',
    'Leche y lactosa','Frutos de cáscara','Apio','Mostaza','Sésamo','Sulfitos','Altramuces','Moluscos']

PLATOS = [
    dict(codigo='PL-001', nombre='Papas arrugadas con mojo rojo', categoria='Entrante', raciones=1, pvp=4.50, extras_pct=6, packaging=0,
         lineas=[('papa',250),('pim_rojo',30),('aceite',15),('ajo',6),('comino',1),('sal',8),('vinagre',5),('pimenton',2)],
         ficha=dict(elaboracion='Lavar las papas sin pelar y cocer en agua muy salada 25-30 min hasta arrugar. Escurrir y secar al fuego. Mojo: triturar pimiento, ajo, comino, pimentón, vinagre, sal y aceite.',
                     conservacion='Mojo en frío positivo (0-4°C) en recipiente cerrado. Papas: elaboración al momento.',
                     temperatura='Servicio caliente, mojo a temperatura ambiente.',
                     vida='Mojo: 5 días refrigerado. Papas: consumo inmediato.',
                     notas='No pelar la papa. Punto de sal alto en cocción, sin sal añadida al emplatar.')),
    dict(codigo='PL-002', nombre='Ropa vieja canaria', categoria='Principal', raciones=1, pvp=9.50, extras_pct=8, packaging=0,
         lineas=[('ternera',160),('garbanzo',120),('papa',120),('cebolla',50),('pim_rojo',40),('tomate',60),('aceite',20),('ajo',5),('vino',30)],
         ficha=dict(elaboracion='Guisar la falda hasta que se deshebre. Sofrito de cebolla, pimiento, ajo y tomate; añadir carne deshebrada, garbanzos y papas fritas en dados. Ligar y reposar.',
                     conservacion='Frío positivo 0-4°C tapado. Regenerar a +65°C en corazón.',
                     temperatura='+65°C.',
                     vida='3 días refrigerado, 2 meses congelado.',
                     notas='Sulfitos por el vino. Mejora al día siguiente.')),
    dict(codigo='PL-003', nombre='Vieja a la espalda con mojo verde', categoria='Principal', raciones=1, pvp=14.00, extras_pct=6, packaging=0,
         lineas=[('vieja',300),('papa',150),('pim_verde',25),('aceite',20),('ajo',6),('sal',5),('vinagre',5)],
         ficha=dict(elaboracion='Abrir la vieja a la espalda, plancha fuerte por piel, terminar con refrito de ajo. Acompañar de papas y mojo verde.',
                     conservacion='Pescado fresco 0-2°C sobre hielo, máx 24h.',
                     temperatura='+63°C.',
                     vida='Consumo en el día.',
                     notas='Merma alta (35%): pedir limpia al proveedor si compensa el sobreprecio.')),
    dict(codigo='PL-004', nombre='Queso asado con miel de palma', categoria='Entrante', raciones=1, pvp=7.00, extras_pct=4, packaging=0,
         lineas=[('queso',140),('miel_palma',15),('gofio',8),('aceite',5)],
         ficha=dict(elaboracion='Marcar el queso a la plancha por ambas caras. Emplatar con miel de palma y espolvoreo de gofio tostado.',
                     conservacion='Queso en frío positivo 0-4°C envuelto.',
                     temperatura='Caliente, servicio inmediato.',
                     vida='Queso abierto: 7 días.',
                     notas='Confirmar con proveedor si el gofio es 100% millo o mezcla con trigo.')),
    dict(codigo='PL-005', nombre='Plátano frito con gofio y miel', categoria='Postre', raciones=1, pvp=5.50, extras_pct=4, packaging=0,
         lineas=[('platano',180),('miel_palma',12),('gofio',10),('aceite',10)],
         ficha=dict(elaboracion='Freír el plátano maduro en aceite a 170°C hasta dorar. Emplatar con gofio y miel de palma.',
                     conservacion='Elaboración al momento.',
                     temperatura='Caliente.',
                     vida='Consumo inmediato.',
                     notas='Usar plátano bien maduro; la merma del 30% corresponde a la piel.')),
]

RESTAURANTE = dict(nombre='La Tasca del Puerto', lugar='Guía de Isora, Tenerife', zona='Canarias', impuesto=7, fc_objetivo=30)

DV_ALERGENO = DataValidation(type="list", formula1='"No contiene,Contiene,Puede contener trazas,Pendiente de validar"', allow_blank=True)


def style_header_row(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = HEAD
        cell.fill = HEAD_FILL
        cell.border = BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def build_plato_sheet(wb, plato):
    nombre_hoja = plato['nombre'][:31]
    ws = wb.create_sheet(nombre_hoja)
    ws.sheet_view.showGridLines = False
    for col, w in zip('ABCDEFG', [26, 12, 9, 15, 10, 15, 13]):
        ws.column_dimensions[col].width = w

    ws['A1'] = f"Escandallo · {plato['nombre']}"
    ws['A1'].font = TITLE
    ws.merge_cells('A1:G1')
    ws['A2'] = f"Código: {plato['codigo']}  ·  Categoría: {plato['categoria']}  ·  Zona: {RESTAURANTE['zona']} (IGIC)"
    ws['A2'].font = Font(italic=True, color="6E6A5C")
    ws.merge_cells('A2:G2')

    r = 4
    headers = ['Ingrediente', 'Cantidad', 'Unidad', 'Precio compra (€/kg-l)', 'Merma %', 'Coste útil (€/kg-l)', 'Coste línea (€)']
    for i, h in enumerate(headers):
        ws.cell(row=r, column=i + 1, value=h)
    style_header_row(ws, r, len(headers))

    first_data_row = r + 1
    row = first_data_row
    for ing_id, cant in plato['lineas']:
        ing = INGREDIENTES[ing_id]
        ws.cell(row=row, column=1, value=ing['nombre'])
        ws.cell(row=row, column=2, value=cant)
        ws.cell(row=row, column=3, value=f"{ing['unidad']} (g/ml)" if ing['unidad'] != 'ud' else 'ud')
        c_precio = ws.cell(row=row, column=4, value=ing['precio']); c_precio.font = BLUE
        c_merma = ws.cell(row=row, column=5, value=ing['merma']); c_merma.font = BLUE
        ws.cell(row=row, column=6, value=f"=D{row}/(1-E{row}/100)").font = BLACK
        ws.cell(row=row, column=7, value=f"=(B{row}/1000)*F{row}").font = BLACK
        for c in range(1, 8):
            ws.cell(row=row, column=c).border = BORDER
        row += 1
    last_data_row = row - 1

    r = row + 1
    labels_formulas = [
        ("Coste ingredientes (€)", f"=SUM(G{first_data_row}:G{last_data_row})", False),
        ("Condimentos y varios (%)", plato['extras_pct'], True),
        ("Extras (€)", None, False),  # formula set below
        ("Packaging (€)", plato['packaging'], True),
        ("Coste directo (€)", None, False),
        ("Raciones", plato['raciones'], True),
        ("Coste por ración (€)", None, False),
    ]
    cell_coste_ing = f"B{r}"
    ws.cell(row=r, column=1, value=labels_formulas[0][0]).font = BOLD
    ws.cell(row=r, column=2, value=labels_formulas[0][1])
    r += 1
    cell_extras_pct = f"B{r}"
    ws.cell(row=r, column=1, value=labels_formulas[1][0])
    c = ws.cell(row=r, column=2, value=labels_formulas[1][1]); c.font = BLUE
    r += 1
    cell_extras = f"B{r}"
    ws.cell(row=r, column=1, value=labels_formulas[2][0])
    ws.cell(row=r, column=2, value=f"={cell_coste_ing}*{cell_extras_pct}/100")
    r += 1
    cell_packaging = f"B{r}"
    ws.cell(row=r, column=1, value=labels_formulas[3][0])
    c = ws.cell(row=r, column=2, value=labels_formulas[3][1]); c.font = BLUE
    r += 1
    cell_coste_directo = f"B{r}"
    ws.cell(row=r, column=1, value=labels_formulas[4][0]).font = BOLD
    ws.cell(row=r, column=2, value=f"={cell_coste_ing}+{cell_extras}+{cell_packaging}").font = BOLD
    r += 1
    cell_raciones = f"B{r}"
    ws.cell(row=r, column=1, value=labels_formulas[5][0])
    c = ws.cell(row=r, column=2, value=labels_formulas[5][1]); c.font = BLUE
    r += 1
    cell_coste_racion = f"B{r}"
    ws.cell(row=r, column=1, value=labels_formulas[6][0]).font = BOLD
    ws.cell(row=r, column=2, value=f"={cell_coste_directo}/{cell_raciones}").font = BOLD
    r += 2

    ws.cell(row=r, column=1, value="PVP carta (€, con IGIC)").font = BOLD
    cell_pvp = f"B{r}"
    c = ws.cell(row=r, column=2, value=plato['pvp']); c.font = BLUE
    r += 1
    ws.cell(row=r, column=1, value="IGIC / IVA (%)")
    cell_igic = f"B{r}"
    c = ws.cell(row=r, column=2, value=RESTAURANTE['impuesto']); c.font = BLUE
    r += 1
    ws.cell(row=r, column=1, value="PVP sin impuesto (€)")
    cell_pvp_base = f"B{r}"
    ws.cell(row=r, column=2, value=f"={cell_pvp}/(1+{cell_igic}/100)")
    r += 1
    ws.cell(row=r, column=1, value="Food cost (%)").font = BOLD
    cell_fc = f"B{r}"
    ws.cell(row=r, column=2, value=f"={cell_coste_racion}/{cell_pvp_base}*100").font = BOLD
    r += 1
    ws.cell(row=r, column=1, value="Margen bruto (€)")
    cell_margen = f"B{r}"
    ws.cell(row=r, column=2, value=f"={cell_pvp_base}-{cell_coste_racion}")
    r += 1
    ws.cell(row=r, column=1, value="Margen (%)")
    cell_margen_pct = f"B{r}"
    ws.cell(row=r, column=2, value=f"={cell_margen}/{cell_pvp_base}*100")
    r += 1
    ws.cell(row=r, column=1, value="Food cost objetivo (%)")
    cell_fc_obj = f"B{r}"
    c = ws.cell(row=r, column=2, value=RESTAURANTE['fc_objetivo']); c.font = BLUE
    r += 1
    ws.cell(row=r, column=1, value="PVP recomendado (€, con IGIC)")
    cell_pvp_rec = f"B{r}"
    ws.cell(row=r, column=2, value=f"=CEILING(({cell_coste_racion}/({cell_fc_obj}/100))*(1+{cell_igic}/100),0.05)")
    r += 1
    ws.cell(row=r, column=1, value="Diagnóstico").font = BOLD
    cell_diag = f"B{r}"
    ws.cell(row=r, column=2,
            value=f'=IF({cell_fc}>{cell_fc_obj}+5,"No rentable",IF({cell_fc}>{cell_fc_obj},"Revisar","Rentable"))').font = BOLD
    r += 2

    ws.cell(row=r, column=1, value="Alérgenos (14 UE)").font = BOLD
    r += 1
    ws.cell(row=r, column=1, value="Alérgeno"); ws.cell(row=r, column=2, value="Estado")
    style_header_row(ws, r, 2)
    r += 1
    alergenos_plato = {}
    for ing_id, _ in plato['lineas']:
        for al, estado in ALERGENOS_POR_INGREDIENTE.get(ing_id, {}).items():
            rank = {'Contiene': 3, 'Pendiente de validar': 2, 'Puede contener trazas': 1}
            if al not in alergenos_plato or rank[estado] > rank[alergenos_plato[al]]:
                alergenos_plato[al] = estado
    for al in LISTA_14_ALERGENOS:
        ws.cell(row=r, column=1, value=al)
        c = ws.cell(row=r, column=2, value=alergenos_plato.get(al, 'No contiene'))
        DV_ALERGENO.add(c)
        r += 1
    r += 1
    ws.cell(row=r, column=1,
            value="Validar siempre con información real de proveedor y etiquetado de producto antes de publicar la carta.").font = Font(italic=True, size=9, color="8A8474")
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=7)
    r += 2

    ws.cell(row=r, column=1, value="Ficha técnica").font = BOLD
    r += 1
    for label, key in [('Elaboración', 'elaboracion'), ('Conservación', 'conservacion'),
                        ('Temperatura de servicio', 'temperatura'), ('Vida útil', 'vida'), ('Observaciones', 'notas')]:
        ws.cell(row=r, column=1, value=label).font = BOLD
        cell = ws.cell(row=r, column=2, value=plato['ficha'][key])
        cell.alignment = WRAP
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=7)
        ws.row_dimensions[r].height = 30
        r += 1

    if DV_ALERGENO not in ws.data_validations.dataValidation:
        ws.add_data_validation(DV_ALERGENO)

    return dict(nombre_hoja=nombre_hoja, coste_racion=cell_coste_racion, pvp=cell_pvp, pvp_base=cell_pvp_base,
                fc=cell_fc, margen=cell_margen, margen_pct=cell_margen_pct, fc_obj=cell_fc_obj,
                pvp_rec=cell_pvp_rec, diag=cell_diag, raciones=cell_raciones)


def build_resumen_sheet(wb, refs):
    ws = wb.create_sheet('Resumen carta', 0)
    ws.sheet_view.showGridLines = False
    ws['A1'] = f"Resumen de carta · {RESTAURANTE['nombre']} · {RESTAURANTE['lugar']}"
    ws['A1'].font = TITLE
    ws.merge_cells('A1:J1')
    ws['A2'] = f"Zona fiscal: {RESTAURANTE['zona']} (IGIC {RESTAURANTE['impuesto']}%)  ·  Food cost objetivo: {RESTAURANTE['fc_objetivo']}%"
    ws['A2'].font = Font(italic=True, color="6E6A5C")
    ws.merge_cells('A2:J2')

    headers = ['Código', 'Plato', 'Categoría', 'Raciones', 'Coste/ración (€)', 'PVP (€)', 'Food cost (%)',
                'Margen bruto (€)', 'Margen (%)', 'PVP recomendado (€)', 'Diagnóstico']
    r = 4
    for i, h in enumerate(headers):
        ws.cell(row=r, column=i + 1, value=h)
    style_header_row(ws, r, len(headers))
    widths = [10, 30, 12, 9, 14, 10, 12, 15, 11, 16, 13]
    for i, w in enumerate(widths):
        ws.column_dimensions[get_column_letter(i + 1)].width = w

    r += 1
    for plato, ref in zip(PLATOS, refs):
        sn = ref['nombre_hoja']
        ws.cell(row=r, column=1, value=plato['codigo'])
        ws.cell(row=r, column=2, value=plato['nombre'])
        ws.cell(row=r, column=3, value=plato['categoria'])
        ws.cell(row=r, column=4, value=f"='{sn}'!{ref['raciones']}")
        ws.cell(row=r, column=5, value=f"='{sn}'!{ref['coste_racion']}")
        ws.cell(row=r, column=6, value=f"='{sn}'!{ref['pvp']}")
        ws.cell(row=r, column=7, value=f"='{sn}'!{ref['fc']}")
        ws.cell(row=r, column=8, value=f"='{sn}'!{ref['margen']}")
        ws.cell(row=r, column=9, value=f"='{sn}'!{ref['margen_pct']}")
        ws.cell(row=r, column=10, value=f"='{sn}'!{ref['pvp_rec']}")
        ws.cell(row=r, column=11, value=f"='{sn}'!{ref['diag']}")
        for c in range(1, 12):
            ws.cell(row=r, column=c).border = BORDER
        r += 1

    last = r - 1
    for col_letter, fmt in [('E', '#,##0.00" €"'), ('F', '#,##0.00" €"'), ('G', '0.0"%"'),
                              ('H', '#,##0.00" €"'), ('I', '0.0"%"'), ('J', '#,##0.00" €"')]:
        for row in range(5, last + 1):
            ws[f"{col_letter}{row}"].number_format = fmt

    r += 1
    ws.cell(row=r, column=1, value="Food cost medio carta").font = BOLD
    ws.cell(row=r, column=7, value=f"=AVERAGE(G5:G{last})").font = BOLD
    ws.cell(row=r, column=7).number_format = '0.0"%"'
    r += 1
    ws.cell(row=r, column=1, value="Platos a revisar / no rentables").font = BOLD
    ws.cell(row=r, column=7, value=f'=COUNTIF(K5:K{last},"Revisar")+COUNTIF(K5:K{last},"No rentable")').font = BOLD

    return ws


def main():
    wb = Workbook()
    wb.remove(wb.active)
    refs = [build_plato_sheet(wb, p) for p in PLATOS]
    build_resumen_sheet(wb, refs)
    wb.save('/home/claude/xlsx_build/chef_online_escandallos.xlsx')


if __name__ == '__main__':
    main()
